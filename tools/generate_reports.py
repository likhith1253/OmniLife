import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client

REPORTS_DIR = r"D:\OmniLife\docs\reports"

FILES = [
    "Likhith_Individual_Contribution_Report",
    "Sunil_Individual_Contribution_Report",
    "Prudhvi_Individual_Contribution_Report",
    "Anish_Individual_Contribution_Report",
]

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def style_heading(p, text, level=1):
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.bold = True
    if level == 1:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(31, 78, 121)
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(46, 117, 182)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(68, 84, 106)
    return p

def add_inline_formatted(paragraph, text, base_font_size=10.5, is_bullet=False):
    # Regex to handle bold, inline code, and normal text
    # e.g. **bold**, `code`, normal
    parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = RGBColor(30, 30, 30)
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 0.5)
            run.font.color.rgb = RGBColor(160, 40, 40)
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = RGBColor(40, 40, 40)

def convert_md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = [line.rstrip('\r\n') for line in f.readlines()]

    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    i = 0
    while i < len(lines):
        line = lines[i]
        
        # Blank line
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Header 1
        if line.startswith('# '):
            p = doc.add_paragraph()
            style_heading(p, line[2:].strip(), level=1)
            i += 1
            continue

        # Header 2
        if line.startswith('## '):
            p = doc.add_paragraph()
            style_heading(p, line[3:].strip(), level=2)
            i += 1
            continue

        # Header 3
        if line.startswith('### '):
            p = doc.add_paragraph()
            style_heading(p, line[4:].strip(), level=3)
            i += 1
            continue

        # Header 4
        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[5:].strip())
            run.font.name = 'Calibri'
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(31, 78, 121)
            i += 1
            continue

        # Markdown Table
        if line.startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < len(lines) and lines[i].startswith('|'):
                # skip separator rows e.g. |---|---|
                if not re.match(r'^\|(\s*[-:]+[-|\s:]*)\|$', lines[i].strip()):
                    table_lines.append(lines[i])
                i += 1
            
            if table_lines:
                rows_data = [[c.strip() for c in t_line.strip('|').split('|')] for t_line in table_lines]
                num_cols = max(len(r) for r in rows_data)
                tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = True

                for r_idx, row in enumerate(rows_data):
                    is_header = (r_idx == 0)
                    for c_idx, cell_value in enumerate(row):
                        if c_idx < num_cols:
                            cell = tbl.cell(r_idx, c_idx)
                            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            if is_header:
                                set_cell_background(cell, "2E75B6")
                                run = p.add_run(cell_value)
                                run.bold = True
                                run.font.name = 'Calibri'
                                run.font.size = Pt(10)
                                run.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "F2F4F7")
                                add_inline_formatted(p, cell_value, base_font_size=9.5)
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # Bullet list items
        if line.startswith('- ') or line.startswith('* ') or re.match(r'^\s{2,4}[-*]\s', line):
            indent_level = 1 if re.match(r'^\s{2,4}[-*]\s', line) else 0
            raw_text = re.sub(r'^\s*[-*]\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            add_inline_formatted(p, raw_text, base_font_size=10.5, is_bullet=True)
            i += 1
            continue

        # Numbered list items
        num_match = re.match(r'^\s*(\d+)\.\s+(.*)', line)
        if num_match:
            raw_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_inline_formatted(p, raw_text, base_font_size=10.5)
            i += 1
            continue

        # Standard paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_inline_formatted(p, line, base_font_size=10.5)
        i += 1

    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")

def convert_all():
    # 1. Generate DOCX
    for name in FILES:
        md_file = os.path.join(REPORTS_DIR, f"{name}.md")
        docx_file = os.path.join(REPORTS_DIR, f"{name}.docx")
        convert_md_to_docx(md_file, docx_file)

    # 2. Convert to PDF using MS Word COM
    print("Converting DOCX to PDF via MS Word...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        for name in FILES:
            docx_file = os.path.abspath(os.path.join(REPORTS_DIR, f"{name}.docx"))
            pdf_file = os.path.abspath(os.path.join(REPORTS_DIR, f"{name}.pdf"))
            print(f"Opening {docx_file}...")
            doc = word.Documents.Open(docx_file)
            # 17 = wdExportFormatPDF
            doc.SaveAs(pdf_file, FileFormat=17)
            doc.Close()
            print(f"Generated PDF: {pdf_file}")
    finally:
        word.Quit()
        print("Done!")

if __name__ == '__main__':
    convert_all()
